#!/usr/bin/env python3
"""Extract top 25 from Arena AI leaderboards.

Usage:
  python sources/leaderboard.py

Description:
  Fetches the Arena AI leaderboard pages
  using curl, extracts the top 25 models
  with their scores, and saves results
  into Excel files (.xlsx) and one JSON
  file.

  Each model name in the Excel output is
  a clickable hyperlink to its Arena page.

  The JSON file is the contract read by
  g4_update_bench.py, which draws the
  Benchmarks slide as two plain columns.
  This module owns vendor classification,
  so the renderer never repeats those rules.

Output files:
  leaderboard_text.xlsx
  leaderboard_text_coding.xlsx
  data/leaderboard.json

Requirements:
  pip install openpyxl

Created: 2026-07-10
Last updated: 2026-09-12
"""

import datetime as dt
import json
import re
import os
import subprocess
import sys
from urllib.parse import quote
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment
)

# work/data, whatever directory the command runs from.
DATA_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "data")

BOARDS = [
    {
        "url": "https://arena.ai/leaderboard/text",
        "label": "English",
        "xlsx": os.path.join(DATA_DIR, "leaderboard_text.xlsx"),
    },
    {
        "url": (
            "https://arena.ai/leaderboard/text/coding"
        ),
        "label": "Coding",
        "xlsx": os.path.join(DATA_DIR,
                             "leaderboard_text_coding.xlsx"),
    },
]

JSON_OUT = os.path.join(DATA_DIR, "leaderboard.json")

NUM_ROWS = 25
FONT_NAME = "Calibri"
FONT_SIZE = 8

VENDOR_RULES = [
    ("anthropic", ["claude"]),
    ("google", [
        "gemini", "gemma", "palm", "bard"
    ]),
    ("openai", [
        "gpt-", "chatgpt", "o1-", "o3-"
    ]),
    ("opensource", [
        "llama", "mistral", "qwen",
        "deepseek", "yi-", "phi-",
        "command-r", "dbrx", "falcon",
        "vicuna", "wizardlm", "solar",
        "mixtral", "olmo", "jamba",
        "internlm", "glm", "kimi",
    ]),
]

VENDOR_COLORS = {
    "anthropic": "4472C4",
    "google":    "FF0000",
    "openai":    "FFD700",
    "opensource": "00B050",
    "other":     "D9D9D9",
}

# --------------------------------------------------------------
def fetch_html(url):
    """Fetch HTML content from a URL."""
    result = subprocess.run(
        ["curl", "-s", url],
        capture_output=True,
        text=True,
        timeout=30,
    )
    if result.returncode != 0:
        print(f"Error fetching {url}")
        sys.exit(1)
    return result.stdout

# --------------------------------------------------------------
def extract_entries(html):
    """Extract leaderboard entries from JSON
    embedded (escaped) in the page HTML."""
    marker = '\\"entries\\":'
    idx = html.find(marker)
    if idx == -1:
        marker = '"entries":'
        idx = html.find(marker)
    if idx == -1:
        return []
    start = html.find('[', idx)
    if start == -1:
        return []
    chunk = html[start:start + 500000]
    chunk = chunk.replace('\\"', '"')
    chunk = chunk.replace('\\\\', '\\')
    decoder = json.JSONDecoder()
    try:
        entries, _ = decoder.raw_decode(chunk)
    except json.JSONDecodeError:
        return []
    return entries

# --------------------------------------------------------------
def extract_cutoff(html):
    """Read the date the leaderboard's votes end on.

    Arena publishes it as voteCutoffISOString next to the
    entries. It is what the board actually reflects, which
    is usually a few days before we download it.
    """
    match = re.search(
        r'voteCutoffISOString\\?":\\?"(\d{4}-\d{2}-\d{2})',
        html
    )
    return match.group(1) if match else None

# --------------------------------------------------------------
def model_page_url(name):
    """Build Arena page URL for a model."""
    return (
        "https://arena.ai/leaderboard/model/"
        + quote(name)
    )

# --------------------------------------------------------------
def get_vendor(name):
    """Determine vendor category from name."""
    lower = name.lower()
    for vendor, keywords in VENDOR_RULES:
        for kw in keywords:
            if kw in lower:
                return vendor
    return "other"

# --------------------------------------------------------------
def make_color_fill(vendor):
    """Create a PatternFill for a vendor."""
    color = VENDOR_COLORS.get(
        vendor, VENDOR_COLORS["other"]
    )
    return PatternFill(
        start_color=color,
        end_color=color,
        fill_type="solid",
    )

# --------------------------------------------------------------
def make_fonts():
    """Build the base, header, and link fonts."""
    return {
        "base": Font(
            name=FONT_NAME, size=FONT_SIZE
        ),
        "header": Font(
            name=FONT_NAME, size=FONT_SIZE,
            bold=True
        ),
        "link": Font(
            name=FONT_NAME, size=FONT_SIZE,
            color="0563C1", underline="single"
        ),
    }

# --------------------------------------------------------------
def write_header(ws, fonts):
    """Write and style the header row."""
    ws["A1"] = "Code"
    ws["B1"] = "Model"
    ws["C1"] = "Score"
    for cell in [ws["A1"], ws["B1"], ws["C1"]]:
        cell.font = fonts["header"]

# --------------------------------------------------------------
def write_row(ws, row, entry, fonts):
    """Write one model row with its vendor swatch."""
    vendor = entry["vendor"]
    cell_a = ws.cell(
        row=row, column=1, value="■"
    )
    cell_a.font = Font(
        name=FONT_NAME, size=FONT_SIZE,
        color=VENDOR_COLORS.get(
            vendor, VENDOR_COLORS["other"]
        ),
    )
    cell_a.fill = make_color_fill(vendor)
    cell_b = ws.cell(
        row=row, column=2, value=entry["name"]
    )
    cell_b.hyperlink = entry["url"]
    cell_b.font = fonts["link"]
    ws.cell(
        row=row, column=3, value=entry["score"]
    ).font = fonts["base"]

# --------------------------------------------------------------
def apply_sheet_layout(ws, count):
    """Set column widths, row heights, alignment."""
    no_pad = Alignment(
        indent=0, wrap_text=False,
        vertical="center"
    )
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 6
    for row in range(1, count + 2):
        ws.row_dimensions[row].height = 11
    for row_cells in ws.iter_rows(
        min_row=1, max_row=count + 1,
        min_col=1, max_col=3
    ):
        for cell in row_cells:
            cell.alignment = no_pad

# --------------------------------------------------------------
def save_to_excel(rows, output_file):
    """Save leaderboard rows to an Excel file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Leaderboard"
    fonts = make_fonts()
    write_header(ws, fonts)
    for index, entry in enumerate(rows):
        write_row(ws, index + 2, entry, fonts)
    apply_sheet_layout(ws, len(rows))
    wb.save(output_file)

# --------------------------------------------------------------
def build_rows(names, urls, scores, num_rows):
    """Build the top-N rows as plain dictionaries."""
    count = min(
        len(names), len(urls),
        len(scores), num_rows
    )
    return [
        {
            "name": names[i],
            "score": int(round(scores[i])),
            "vendor": get_vendor(names[i]),
            "url": urls[i],
        }
        for i in range(count)
    ]

# --------------------------------------------------------------
def process_board(board):
    """Fetch one board, save XLSX, return its rows."""
    url = board["url"]
    print(f"Fetching {url} ...")
    html = fetch_html(url)
    entries = extract_entries(html)
    entries.sort(
        key=lambda e: e.get("rank", 10**9)
    )
    names = [
        e["modelDisplayName"] for e in entries
    ]
    model_urls = [
        model_page_url(n) for n in names
    ]
    scores = [
        e["rating"] for e in entries
    ]
    print(f"  Found {len(names)} models")
    if not names:
        print(
            "  ERROR: no models found -- "
            "page format may have changed"
        )
        sys.exit(1)
    rows = build_rows(
        names, model_urls, scores, NUM_ROWS
    )
    save_to_excel(rows, board["xlsx"])
    print(f"  Saved to {board['xlsx']}")
    return rows, extract_cutoff(html)

# --------------------------------------------------------------
def save_to_json(boards, output_file):
    """Write the boards as one JSON contract file."""
    directory = os.path.dirname(output_file)
    if directory:
        os.makedirs(directory, exist_ok=True)
    payload = {
        "fetched": dt.date.today().isoformat(),
        "boards": boards,
    }
    with open(output_file, "w",
              encoding="utf-8") as handle:
        json.dump(payload, handle, indent=2)
    print(f"  Saved to {output_file}")

# --------------------------------------------------------------
def main():
    """Main function."""
    boards = []
    for board in BOARDS:
        rows, cutoff = process_board(board)
        print(f"  Votes counted through {cutoff}")
        boards.append({
            "label": board["label"],
            "url": board["url"],
            "cutoff": cutoff,
            "entries": rows,
        })
    save_to_json(boards, JSON_OUT)
    print("\nDone!")

# --------------------------------------------------------------
if __name__ == "__main__":
    main()
